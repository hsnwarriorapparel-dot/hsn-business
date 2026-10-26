#!/usr/bin/env python3
"""Daily 2PM PKT lead scraper: Google Maps engine (with email extraction) + website enrichment.
Feeds leads.csv for the 1AM outreach pipeline. Rotates 2 cities/day."""
import json, csv, os, re, time, tempfile, urllib.request, datetime, subprocess, sys, shutil

GMAPS = "http://localhost:8080/api/v1/jobs"
LEADS = "/home/ubuntu/leads.csv"
LOG = "/home/ubuntu/gmaps_scrape_log.txt"
DATE = datetime.date.today().isoformat()

# City rotation (lat, lon, country, phone intl prefix)
CITIES = [
    ("Houston TX", "29.7604", "-95.3698", "USA", "+1"),
    ("Chicago IL", "41.8781", "-87.6298", "USA", "+1"),
    ("Phoenix AZ", "33.4484", "-112.0740", "USA", "+1"),
    ("Philadelphia PA", "39.9526", "-75.1652", "USA", "+1"),
    ("San Antonio TX", "29.4241", "-98.4936", "USA", "+1"),
    ("Dallas TX", "32.7767", "-96.7970", "USA", "+1"),
    ("Atlanta GA", "33.7490", "-84.3880", "USA", "+1"),
    ("Boston MA", "42.3601", "-71.0589", "USA", "+1"),
    ("Seattle WA", "47.6062", "-122.3321", "USA", "+1"),
    ("Denver CO", "39.7392", "-104.9903", "USA", "+1"),
    ("Las Vegas NV", "36.1699", "-115.1398", "USA", "+1"),
    ("Detroit MI", "42.3314", "-83.0458", "USA", "+1"),
    ("London", "51.5072", "-0.1276", "UK", "+44"),
    ("Manchester", "53.4808", "-2.2426", "UK", "+44"),
    ("Dublin", "53.3498", "-6.2603", "Ireland", "+353"),
    ("Amsterdam", "52.3676", "4.9041", "Netherlands", "+31"),
    ("Paris", "48.8566", "2.3522", "France", "+33"),
    ("Berlin", "52.5200", "13.4050", "Germany", "+49"),
    ("Madrid", "40.4168", "-3.7038", "Spain", "+34"),
    ("Barcelona", "41.3874", "2.1686", "Spain", "+34"),
    ("Milan", "45.4642", "9.1900", "Italy", "+39"),
    ("Rome", "41.9028", "12.4964", "Italy", "+39"),
    ("Lisbon", "38.7223", "-9.1393", "Portugal", "+351"),
    ("Zurich", "47.3769", "8.5417", "Switzerland", "+41"),
]

KEYWORD_TYPES = ["boxing gym", "MMA gym", "jiu jitsu academy", "muay thai gym", "martial arts school"]

BLOCKED_EMAIL = ['sentry', 'wixpress', 'example.com', 'schema', 'abuse@', 'postmaster@', 'privacy@',
                 'godaddy', 'domain', 'youremail', 'email@', 'name@', 'test@']

WEBMAIL_DOMAINS = {'gmail.com', 'hotmail.com', 'outlook.com', 'yahoo.com', 'yahoo.fr', 'yahoo.it',
                    'yahoo.es', 'yahoo.de', 'live.com', 'msn.com', 'aol.com', 'mail.com', 'gmx.de',
                    'gmx.net', 'gmx.ch', 'web.de', 't-online.de', 'orange.fr', 'free.fr', 'laposte.net',
                    'libero.it', 'virgilio.it', 'tiscali.it', 'protonmail.com', 'proton.me',
                    'icloud.com', 'me.com', 'zoho.com', 'bluewin.ch'}

def log(m):
    line = time.strftime('%H:%M:%S') + ' ' + m
    print(line, flush=True)
    with open(LOG, 'a') as f:
        f.write(line + '\n')

def api(url, data=None, timeout=20):
    req = urllib.request.Request(url, headers={"Content-Type": "application/json"})
    if data is not None:
        req.data = json.dumps(data).encode()
    with urllib.request.urlopen(req, timeout=timeout) as r:
        return json.loads(r.read().decode())

def engine_up():
    try:
        urllib.request.urlopen(GMAPS, timeout=5)
        return True
    except Exception:
        return False

def ensure_engine():
    if not engine_up():
        log("engine down - starting docker container")
        subprocess.run(["sudo", "docker", "compose", "up", "-d"],
                       cwd="/home/ubuntu/skill-installs/google-maps-scraper-kit",
                       capture_output=True, timeout=120)
        time.sleep(15)
        if not engine_up():
            log("ENGINE FAILED TO START")
            sys.exit(1)

def run_city(city, lat, lon, country, prefix, max_time=900):
    """One gmaps job per city; returns list of lead dicts."""
    name = f"gmaps-{city.lower().replace(' ', '-')}-{DATE}"
    payload = {
        "name": name,
        "keywords": [f"{kt} in {city}" for kt in KEYWORD_TYPES],
        "lat": str(lat), "lon": str(lon), "lang": "en", "depth": 5,
        "max_time": max_time, "email": True,
    }
    j = api(GMAPS, payload)
    jid = j["id"]
    log(f"job created for {city}: {jid}")
    # poll
    deadline = time.time() + max_time + 300
    while time.time() < deadline:
        time.sleep(30)
        jobs = api(GMAPS, timeout=30)
        job = next((x for x in jobs if x["ID"] == jid), None)
        st = job["Status"] if job else "missing"
        if st in ("ok", "done", "complete"):
            log(f"{city}: job done")
            break
        log(f"{city}: {st}...")
    # download
    csv_url = f"{GMAPS}/{jid}/download"
    try:
        with urllib.request.urlopen(csv_url, timeout=60) as r:
            raw = r.read().decode('utf-8', errors='ignore')
    except Exception as e:
        log(f"{city}: download failed {str(e)[:50]}")
        return []
    rows = list(csv.reader(raw.splitlines()))
    if len(rows) < 2:
        log(f"{city}: no results")
        return []
    hdr = rows[0]
    i = {h: k for k, h in enumerate(hdr)}
    leads = []
    for r in rows[1:]:
        if len(r) < 5:
            continue
        def g(col):
            return r[i[col]].strip() if col in i and len(r) > i[col] else ''
        title = g('title')
        if not title:
            continue
        leads.append({
            "name": title, "category": g('category')[:60], "phone_raw": g('phone'),
            "website": g('website'), "email": (g('emails') or g('email')).lower(), "address": g('address')[:100],
            "city": city, "country": country, "prefix": prefix,
        })
    log(f"{city}: {len(leads)} leads scraped")
    return leads

def fetch(url, timeout=10):
    try:
        req = urllib.request.Request(url, headers={'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36'})
        with urllib.request.urlopen(req, timeout=timeout) as r:
            return r.read().decode('utf-8', errors='ignore')
    except Exception:
        return ''

def extract_emails(html):
    found = re.findall(r'[a-zA-Z0-9._%+-]+@[a-zA-Z0-9.-]+\.[a-zA-Z]{2,}', html or '')
    return [e.lower() for e in found if not any(b in e.lower() for b in BLOCKED_EMAIL)]

def normalize_phone(raw, prefix):
    digits = re.sub(r'\D', '', raw or '')
    if not digits:
        return ''
    if prefix == '+1' and len(digits) == 10:
        return '+1' + digits
    if prefix == '+1' and len(digits) == 11 and digits.startswith('1'):
        return '+' + digits
    if digits.startswith('0'):
        return ''
    if not raw.strip().startswith('+') and prefix and len(digits) >= 9:
        return prefix + digits.lstrip('0')
    if raw.strip().startswith('+'):
        return '+' + digits
    return digits

def enrich_website(lead):
    """If gmaps didn't get an email, fetch the website + contact pages."""
    site = lead["website"]
    if not site or lead["email"]:
        return lead["email"]
    if not site.startswith('http'):
        site = 'https://' + site
    pages = [site.rstrip('/') + p for p in ['', '/contact', '/contact-us', '/kontakt', '/impressum']]
    for p in pages[:4]:
        html = fetch(p)
        if html:
            ems = extract_emails(html)
            if ems:
                lead["email"] = ems[0]
                break
        time.sleep(0.5)
    return lead["email"]

def clean_email(cell, website='', name=''):
    """The engine's `emails` column can be a comma-separated list mixing a real
    address with junk (image-filename pseudo-addresses, placeholder domains,
    third-party platform addresses). Store ONE clean address or '' so the row
    is not auto-excluded later. Rule (mirrors repair_tighten.py):
      - reject image/placeholder/hash/third-party/PEC candidates
      - accept an exact or base-domain match with the lead's own website
      - accept a webmail address whose local part matches the lead name/site
      - accept a business domain whose slug matches the lead name/site slug
      - otherwise '' (the lead is still stored with its phone/website)
    """
    cands = [c.strip().lower() for c in (cell or '').split(',') if c.strip()]
    if not cands:
        return ''
    cands = [re.sub(r'^(u003e|%20|u00a0|\s)+', '', c).lstrip('-.') for c in cands]
    site = (website or '').strip()
    m = re.match(r'^https?://([^/]+)', site)
    sd = m.group(1).lower() if m else ''
    if sd.startswith('www.'):
        sd = sd[4:]
    sbase = '.'.join(sd.split('.')[-2:]) if sd.count('.') >= 1 else sd
    toks = set()
    for src in (name or '', site):
        for w in re.split(r'[^a-z0-9]+', src.lower()):
            if len(w) >= 4 and w not in ('gym', 'club', 'academy', 'school', 'team',
                                         'martial', 'arts', 'boxing', 'jiu', 'jitsu',
                                         'karate', 'fitness', 'sports', 'center', 'centre'):
                toks.add(w)
    best = ''
    best_score = 0
    for c in cands:
        if not re.match(r'^[a-z0-9._%+\-]+@[a-z0-9.\-]+\.[a-z]{2,}$', c):
            continue
        dom = c.split('@')[1]
        local = c.split('@')[0]
        if re.search(r'\.(jpg|jpeg|png|gif|webp|svg|pdf|html?|css|js)$', dom) or '2x' in dom:
            continue
        if any(b in dom for b in ('sentry', 'wixpress', 'domain.com', 'dominio.com',
                                  'dominio.pt', 'mysite.com', 'example.com', 'exemplo',
                                  'beispiel', 'email.com', 'framer', 'webador')):
            continue
        if any(t in dom for t in ('schedulista', 'ayudapanel', 'fitkit', 'bookizon',
                                   'squarespace', 'wixsite', 'facebook.com', 'instagram.com',
                                   'calendly', 'eventbrite', 'mindbody', 'gymdesk',
                                   'wordpress', 'blogspot', 'weebly', 'jimdo', 'tripadvisor')):
            continue
        if re.fullmatch(r'[0-9a-f]{16,}', local) or re.fullmatch(r'[0-9a-f]{8}-[0-9a-f\-]{8,}', local):
            continue
        if dom.endswith('pec.it'):
            continue
        # score: exact site domain > base-domain match > name-matched webmail
        # > business-domain slug match. Pick the BEST candidate, not the first.
        score = 0
        if sd and dom == sd:
            score = 100
        elif sd and (base_domain(dom) == sbase or dom in sd or sd in dom):
            score = 90
        elif dom in WEBMAIL_DOMAINS:
            for t in toks:
                if len(t) >= 4 and (t in local or (len(local) >= 4 and local in t)):
                    score = max(score, 60)
                    break
        else:
            dslug = dom.split('.')[0]
            for t in toks:
                if len(t) >= 5 and (t in dslug or dslug in t):
                    score = max(score, 50)
                    break
        if score > best_score:
            best, best_score = c, score
    return best


def base_domain(dom):
    p = dom.split('.')
    return '.'.join(p[-2:]) if len(p) >= 2 else dom


def ig_from_website(site):
    m = re.search(r'instagram\.com/([A-Za-z0-9_.]{2,40})', site or '')
    if m:
        h = m.group(1).rstrip('.')
        if h.lower() not in ('p', 'explore', 'accounts', 'stories'):
            return h
    return ''

def append_leads(leads):
    """Dedupe + append to leads.csv. Returns count added."""
    with open(LEADS) as f:
        rows = list(csv.reader(f))
    hdr = rows[0]
    idx = {h: k for k, h in enumerate(hdr)}
    def col(name):
        return idx.get(name)
    existing_emails = {(r[col('Email')] or '').strip().lower() for r in rows[1:] if col('Email') is not None}
    existing_phones = {re.sub(r'\D', '', r[col('Phone')] or '') for r in rows[1:] if col('Phone') is not None}
    # next Lead ID = max existing numeric ID + 1 (row count collides with re-used IDs)
    import re as _re
    max_id = max((int(m.group(1)) for r in rows[1:] if (m := _re.match(r'^L(\d+)$', (r[0] or '').strip()))), default=0)
    next_id = max_id + 1
    added = 0
    for lead in leads:
        em = clean_email(lead["email"], lead.get("website", ""), lead.get("name", ""))
        ph = normalize_phone(lead["phone_raw"], lead["prefix"])
        if em and em in existing_emails:
            continue
        if ph and re.sub(r'\D', '', ph) in existing_phones:
            continue
        if not em and not ph and not lead["website"]:
            continue
        num = next_id
        next_id += 1
        row = [''] * len(hdr)
        row[0] = f"L{num:03d}"
        for field, colname in [("name", "Name"), ("city", "City"), ("country", "Country"),
                               ("category", "Type"), ("website", "Website")]:
            c = col(colname)
            if c is not None:
                row[c] = lead[field][:60]
        c = col('Email')
        if c is not None: row[c] = em
        c = col('Phone')
        if c is not None: row[c] = ph
        ig = ig_from_website(lead["website"])
        c = col('IG Username')
        if c is not None and ig: row[c] = ig
        c = col('Status')
        if c is not None: row[c] = f"HUNTED (gmaps 2PM {DATE})"
        rows.append(row)
        if em: existing_emails.add(em)
        if ph: existing_phones.add(re.sub(r'\D', '', ph))
        added += 1
    fd, tmp = tempfile.mkstemp(dir='/home/ubuntu', suffix='.csv.tmp')
    with os.fdopen(fd, 'w', newline='') as f:
        csv.writer(f).writerows(rows)
    with open(tmp) as f:
        assert len(list(csv.reader(f))) == len(rows)
    os.replace(tmp, LEADS)
    return added

def rebuild_excel():
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill
        with open(LEADS) as f:
            rows = list(csv.reader(f))
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Leads"
        for row in rows:
            ws.append(row)
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F4E78")
        for col in "ABCDEFGHIJKLMNO":
            ws.column_dimensions[col].width = 24
        today = sum(1 for r in rows[1:] if f"gmaps 2PM {DATE}" in (r[-4] if len(r) > 4 else ''))
        ws2 = wb.create_sheet("Scrape Report")
        ws2.append([f"gmaps daily scrape - {DATE}"])
        ws2.append(["New leads today", today])
        ws2.append(["Total", len(rows) - 1])
        wb.save('/home/ubuntu/HWA_Lead_Tracker.xlsx')
        return today
    except Exception as e:
        log(f"excel rebuild failed: {str(e)[:60]}")
        return -1

def git_commit():
    # Refresh backup copies FIRST — business-backup/leads.csv is a separate file, not a
    # symlink, so git add -A alone commits nothing (silent no-op push).
    for src in ("/home/ubuntu/leads.csv", "/home/ubuntu/HWA_Lead_Tracker.xlsx"):
        try:
            shutil.copy2(src, os.path.join("/home/ubuntu/business-backup", os.path.basename(src)))
        except Exception as e:
            log(f"backup copy failed {os.path.basename(src)}: {str(e)[:60]}")
    subprocess.run(["git", "add", "-A"], cwd="/home/ubuntu/business-backup", capture_output=True)
    c = subprocess.run(["git", "commit", "-m", f"gmaps 2PM scrape {DATE}"],
                       cwd="/home/ubuntu/business-backup", capture_output=True, text=True)
    if c.returncode == 0:
        log(f"git committed: {c.stdout.strip().splitlines()[0] if c.stdout.strip() else 'ok'}")
    else:
        log(f"git commit skipped/failed: {(c.stdout + c.stderr).strip()[:120]}")
    p = subprocess.run(["bash", "-c",
                        "GIT_SSH_COMMAND='ssh -i /home/ubuntu/.ssh/hsn_backup_key -o IdentitiesOnly=yes' git push origin main"],
                       cwd="/home/ubuntu/business-backup", capture_output=True, text=True)
    if p.returncode == 0:
        log("git push OK")
    else:
        log(f"git push FAILED: {(p.stdout + p.stderr).strip()[:200]}")

def main():
    log(f"=== GMAPS 2PM SCRAPE START {DATE} ===")
    ensure_engine()
    doy = datetime.date.today().timetuple().tm_yday
    picks = [CITIES[(doy * 2) % len(CITIES)], CITIES[(doy * 2 + 1) % len(CITIES)]]
    log(f"cities today: {picks[0][0]}, {picks[1][0]}")
    all_leads = []
    for city, lat, lon, country, prefix in picks:
        leads = run_city(city, lat, lon, country, prefix)
        # website enrichment for those missing emails
        no_email = [l for l in leads if not l["email"] and l["website"]]
        log(f"{city}: enriching {len(no_email)} websites for emails")
        for l in no_email:
            enrich_website(l)
        got_email = sum(1 for l in leads if l["email"])
        log(f"{city}: {got_email}/{len(leads)} have emails after enrichment")
        all_leads.extend(leads)
    added = append_leads(all_leads)
    log(f"NEW LEADS APPENDED: {added}")
    rebuild_excel()
    git_commit()
    log("=== GMAPS 2PM SCRAPE COMPLETE ===")

if __name__ == "__main__":
    main()
