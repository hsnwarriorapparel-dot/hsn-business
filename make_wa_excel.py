#!/usr/bin/env python3
"""Builds the WhatsApp-pending Excel: all leads with WA numbers not yet WA-messaged.
For when the ban lifts - user asked for a separate sheet."""
import csv, openpyxl
from openpyxl.styles import Font, PatternFill
import sys
sys.path.insert(0, '/home/ubuntu')

with open('/home/ubuntu/leads.csv') as f:
    rows = list(csv.reader(f))
idx = {h: i for i, h in enumerate(rows[0])}

wa_pending = []
for r in rows[1:]:
    phone = (r[idx['Phone']] or '').strip()
    status = r[idx['Status']]
    if phone and len(phone) >= 11 and 'WhatsApp sent' not in status:
        wa_pending.append(r)

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "WA Pending"
ws.append(["Lead ID", "Name", "Country", "Phone", "Email", "IG Handle", "Status", "Notes"])
for r in wa_pending:
    ws.append([r[0], r[1], r[idx['Country']], r[idx['Phone']], r[idx['Email']],
               r[idx['IG Username']], r[idx['Status']], r[idx['Notes']]])
for cell in ws[1]:
    cell.font = Font(bold=True, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor="C00000")
widths = {"A":8, "B":30, "C":12, "D":16, "E":34, "F":28, "G":44, "H":30}
for col, wdt in widths.items():
    ws.column_dimensions[col].width = wdt
wb.save('/home/ubuntu/WA_Pending_List.xlsx')
print(f"WA_Pending_List.xlsx: {len(wa_pending)} leads with WA numbers, not yet messaged")
